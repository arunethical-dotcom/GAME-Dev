"""
One-sheet, copy-and-generate prompt list for the castle assets.
  python3 scripts/tools/make_simple_prompts.py
Writes docs/prompts/CASTLE_PROMPTS.xlsx (one sheet) and docs/prompts/CASTLE_PROMPTS.txt.

Design rules (to avoid the defects seen in the earlier asset packs):
- ONE object per image (no sheets of many pieces, no animation strips: animation is done in code).
- Every prompt is complete on its own: view, details, exact colours, background, shape, and a DO-NOT list.
- Sprites on flat magenta #FF00FF; seamless textures fill the whole image; strips are a band across magenta.
- Sizes are handled by Claude after generation (downscale + pixel snap), so prompts give only the image shape.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
OUT = os.path.join(ROOT, 'docs', 'prompts'); os.makedirs(OUT, exist_ok=True)

# ---------------------------------------------------------------- colours (hex, from the castle plan)
STONE = "grey-blue stone: #151a30 outlines, #3c3e52 shadow, #595768, #706b77 mid, #887f86, #a69a96 lit faces, #dccabb top highlights, mortar #2a2b3c"
WET = "wet stone #2a3040 and #3a4252, dark green algae #1f3326 and #2f4a2e, wet shine #8fb6c8"
ROOF = "blue slate roof: #1f2f5a shadow, #2e4677, #486f91 mid, #6b8fb3, #8fb0d0 highlight; gold trim #8a6420, #c9952e, #e0b040, #ffe08a"
BANNER = "crimson cloth #5a1026, #7e1a39, #b64e3f; gold #c9952e, #e0b040, #ffe08a; iron rod #1e2230, #4a4f60"
IRON = "dark iron #12141c, #23283a, #3a4155, #5b6378"
GLOW = "glowing windows #b0602a, #ffb347, #ffd27a, #fff3c4; stained glass #1d4fa0, #b3262e, #e0b040, #3f8f4f"
WOOD = "dark oak #2e1c12, #4a2e1e, #6b4228, #8a5a36"
FIRE = "fire #7a1a08, #c8380e, #f07a1a, #ffb52e, #ffe680, #fffbe0"
WATER = "water #022448, #033361, #045282, #1d7fb8, #5fc6ff; foam #cfeeff, #ffffff"
GRASS = "grass #24402a, #335a33, #4f7f3f, #6f9a45, #95b65a"
ROCK = "rock #1c2233, #2f3648, #4a5263, #6b7385, #9aa1ae; moss #3a5a2e, #5c7f3a"
FLAG = "pale stone #5d5866, #8c8088, #a89a98, #bfb1aa, #d8cbc2; joints #2a2a3a"
COBBLE = "cobbles #6a6372, #958a8e, #ad9f9d, #b5a49f, #c0b2ac, #dacdc4; joints #2e2e3e"
DIRT = "earth #6b4f36, #8f6e4e, #b08c6a, #c2a48a, #d6bc9e; pebbles #7a7470, #a39a90"
CARPET = "wine-red carpet #4a0e22, #7e1a39, #a3264a with gold edge #c9952e, #e0b040"

# ---------------------------------------------------------------- shared blocks (inlined into every prompt)
STYLE = ("STYLE: crisp high-detail pixel art for a 2D fantasy RPG (modern HD pixel art), clean hard square pixels, dark outlines, "
         "no blur, no soft brush, no painterly look, no 3D render. Warm late-afternoon light from the UPPER LEFT: the left side is lit, the right side is in shade.")
VIEW = ("VIEW: the classic top-down RPG camera: we face the FRONT of the object straight-on, looking down slightly from above, so the front face and a little of the top are visible. "
        "The camera is exactly centred and level left-to-right. The object stands perfectly upright: vertical edges perfectly vertical, horizontal edges perfectly horizontal, "
        "parallel lines stay parallel (no perspective, no vanishing point). Not tilted, not rotated, not isometric, no angled side walls showing.")
VIEW_SIDE = ("VIEW: plain side view, the body level and horizontal, not tilted, not rotated.")
BG = ("BACKGROUND: flat solid pure magenta #FF00FF everywhere around the object, one even colour, no gradient, no floor, no ground, no shadow, no water. "
      "Leave a magenta margin on all four sides; the object is fully inside the image and centred, its base near the bottom. "
      "The object's edges go straight to magenta: no glow, halo, white outline or soft fade.")
TILE = ("LAYOUT: a seamless repeating texture that fills the ENTIRE image edge to edge, seen flat and straight-on (no perspective, no tilt). "
        "Left edge matches right edge and top edge matches bottom edge so copies join with no visible seam. "
        "Even lighting across the whole image (no dark corners, no vignette, no spotlight). No background, no border, no magenta.")
BAND = ("LAYOUT: one long horizontal strip running straight across the FULL width of the image through the middle, cut off by the left and right edges "
        "(the pattern continues past both edges and repeats evenly so it can tile left-right). Flat solid magenta #FF00FF above and below the strip. "
        "The strip is perfectly horizontal and straight-on: not tilted, not curved, no perspective.")
DONT_SPRITE = ("DO NOT include: people, characters, an adventurer or any figure for scale; text, labels, numbers, arrows, grids or rulers; "
               "a second object, copies or variations; {GROUND}shadows on the background; "
               "pink, magenta or purple colours inside the object; any tilt, rotation, isometric or angled view; anything cut off by the image edge.")
DONT_TILE = ("DO NOT include: people, text, labels, borders, frames, a single object in the middle, magenta, vignette or dark edges, "
             "perspective or tilt, a visible seam, an obvious repeating stamp pattern.")

import csv as _csv
_CELLS = {}
for _r in _csv.DictReader(open(os.path.join(ROOT, 'docs/asset_map/asset_placements.csv'))):
    _CELLS.setdefault(_r['id'], [])
    _c = (_r['grid_cell'], int(_r['plan_x']))
    if _c[0] not in [c for c, _ in _CELLS[_r['id']]]: _CELLS[_r['id']].append(_c)
def _ck(c):
    L = c.rstrip('0123456789'); n = 0
    for ch in L: n = n * 26 + ord(ch) - 64
    return (int(c[len(L):]), n)
def ref_ids(key):
    if key in _CELLS: return key
    if key[:-1] in _CELLS: return key[:-1]
    return None
STYLE_ONLY = {'F7B', 'F7D'}
def cells_for(key):
    if key in STYLE_ONLY: return []
    k = ref_ids(key)
    if not k: return []
    west = sorted([c for c, x in _CELLS[k] if x <= 768], key=_ck); east = sorted([c for c, x in _CELLS[k] if x > 768], key=_ck)
    return west + east
REF_HEAD = ("REFERENCE IMAGES (3 images are attached with this prompt): "
            "Image 1 = the castle plan with a grid (columns A to AV from left to right, rows 1 to 32 from top to bottom; each cell's name is written faintly inside it). "
            "Image 2 = the colour legend for the labels. "
            "Image 3 = the same castle plan with asset labels (small dark boxes with codes like TW2, G1, B3; a dot marks the exact spot).")
REF_RULES = ("HOW TO USE THE REFERENCE IMAGES: use them ONLY to copy the design of this one object (its shape, proportions, details, ornaments and colours). "
             "Then draw a completely NEW, clean image of that object alone, following every rule below. "
             "Do NOT copy the grid lines, cell names, labels, label boxes, dots or legend. "
             "Do NOT copy the camera angle of the plan: the plan is an angled overview, but this asset must be redrawn exactly as described in VIEW. "
             "Do NOT include any other part of the castle, or any trees, water, grass or ground from the plan. Do NOT redraw the whole castle.")
def ref_block(key, kind):
    cs = cells_for(key)
    if not cs:
        where = ("This exact object does not appear in the plan: use Images 1 and 3 only to match the style, stone colour and level of detail of the castle.")
    else:
        main = cs[0]; more = cs[1:4]
        also = f" (the same object also appears at {', '.join(more)})" if more else ""
        code = ref_ids(key)
        if kind == 'tile':
            what = "the surface/material at that spot"
        elif kind == 'band':
            what = "that part of the wall at that spot"
        else:
            what = "the object at that spot"
        where = (f"WHERE TO LOOK: in Image 3 find the label that contains the code \"{code}\" at grid cell {main}{also}. "
                 f"Then look at the same cell {main} in Image 1 (no labels covering it) and study {what}.")
    return REF_HEAD + "\n\n" + where + "\n\n" + REF_RULES

def shape_of(w, h):
    r = w / h
    opts = [('1:1 square', 1), ('2:3 tall', 2/3), ('1:2 tall', 0.5), ('9:16 tall', 9/16), ('3:4 tall', 0.75), ('4:3 wide', 4/3),
            ('3:2 wide', 1.5), ('16:9 wide', 16/9), ('2:1 wide', 2), ('21:9 wide', 21/9)]
    return min(opts, key=lambda o: abs(o[1] - r))[0]

# (save_as, title, kind, shape(w,h) or None, subject, details, colours)
L = []
def add(*r): L.append(r)

# ---- 1. wall kit
add('T1A_wall_face_a.png', 'Wall face texture A', 'tile', (2, 1), "castle wall stone texture",
    "Dressed rectangular ashlar blocks laid in perfectly horizontal rows (running bond, every row offset by half a block), each block about twice as wide as tall, about 10 rows from top to bottom. Thin dark mortar lines. Each block has a tiny lit bevel on its top-left edge. A few small chips and hairline cracks. No windows, no doors, no moss.", STONE)
add('T1B_wall_face_b.png', 'Wall face texture B', 'tile', (2, 1), "castle wall stone texture (second variant)",
    "Same stone, same block size and row height as a matching ashlar wall, but a different arrangement of blocks, a few slightly darker weathered blocks, very slight moss in two or three mortar joints. Perfectly horizontal rows.", STONE)
add('T2_wall_waterline.png', 'Wall waterline band', 'band', None, "the bottom band of a castle wall where it stands in a moat",
    "Ashlar stone blocks, darker and wet, a band of dark green-black algae along the lower half, thin vertical drip streaks, a thin bright wet line along the very bottom edge of the strip.", STONE + "; " + WET)
add('T3_merlons_front.png', 'Battlements (front)', 'band', None, "a crenellated castle battlement (merlons) seen from the front",
    "A row of identical square stone merlons with equal gaps between them, about 12 merlons across the image, all exactly the same height and evenly spaced, each with a lit top face and a front face, standing on a continuous stone course. The top edges of all merlons are on one straight horizontal line.", STONE)
add('T6_corbels.png', 'Corbel band', 'band', None, "a band of stepped stone corbels (machicolations) that hangs under castle battlements",
    "A row of identical three-step stacked stone brackets, evenly spaced, about 12 across the image, dark shadow gaps between them, hanging from a continuous stone course along the top of the strip.", STONE)
add('T7_wall_walk.png', 'Wall-walk paving', 'tile', (1, 1), "worn stone paving of a castle wall-walk, seen from directly above",
    "Rectangular flagstone slabs in rows, thin dark joints, gentle wear and a few cracks. Lighter than the castle wall.", FLAG)
add('T4_merlons_back.png', 'Battlements (back side)', 'band', None, "a crenellated castle battlement seen from the INSIDE of the castle",
    "A row of identical square stone merlons, evenly spaced, about 12 across, all in soft shade (we are looking at their back faces), standing on the stone walkway edge. Top edges on one straight line.", STONE)

# ---- 2. towers
TOWER = ("Round castle tower BODY only, with a flat cut top (no roof, no battlements on top). Stone blocks in horizontal rows that curve gently with the round shape. "
         "A ring of stepped stone brackets just below the top. A slightly wider stepped base at the bottom.")
add('TW2_tower_body.png', 'Tower body (standard)', 'sprite', (1, 2), "a round stone castle tower body",
    TOWER + " Two narrow arrow slits stacked in the middle. A small iron rod near the top for hanging a banner. Symmetric left and right.", STONE + "; " + IRON + "; arrow slits #0c0e16")
add('TW3_tower_body_large.png', 'Tower body (large)', 'sprite', (1, 2), "a large round stone castle tower body",
    TOWER + " Two narrow arrow slits and one arched window glowing warm yellow, a carved stone band halfway up. Symmetric left and right.", STONE + "; " + GLOW + "; arrow slits #0c0e16")
add('TW6_tower_roof.png', 'Tower cone roof', 'sprite', (1, 1), "a conical tower roof",
    "A steep 8-sided cone roof covered in blue slate tiles, gold trim along each ridge, a gold-trimmed overhanging rim at the bottom, a gold spike finial with a small orb at the tip. Only the roof, nothing under it. The tip is exactly above the centre. Perfectly symmetric.", ROOF)
add('TW4_tower_base_water.png', 'Tower base in water', 'sprite', (3, 1), "the bottom section of a round stone tower where it rises out of water",
    "A short wide cylinder section of stone blocks, wet and dark, a band of green-black algae, a bright wet line along its bottom edge. The bottom edge is a straight horizontal line (no water drawn). Symmetric.", STONE + "; " + WET)
add('TW5_tower_top_battlements.png', 'Tower top with battlements', 'sprite', (2, 1), "the open top of a round stone tower with battlements, no roof",
    "A ring of 7 square stone merlons around a round flagstone floor that is visible inside, seen from the front and slightly above. Only this top section. Symmetric.", STONE + "; " + FLAG)
add('TW7_square_pillar.png', 'Square pillar tower', 'sprite', (1, 2), "a square stone pedestal tower",
    "A square stone pillar-tower, flat top with a low stone ledge (a statue will stand on it), one recessed niche with a small crimson banner, a stepped base. Symmetric.", STONE + "; " + BANNER)
add('TW8_tower_banner.png', 'Tower banner', 'sprite', (1, 2), "a hanging castle banner",
    "A tall crimson cloth banner hanging straight down from a horizontal iron rod with small gold end caps, a gold cross emblem in the middle, a thin gold border, a swallow-tail cut at the bottom. Hanging still, perfectly straight and symmetric.", BANNER)
add('TW9_pennant.png', 'Pennant flag', 'sprite', (1, 1), "a pennant flag on a pole tip",
    "The top of a thin vertical golden flagpole with a small orb, and a crimson swallow-tail pennant with a small gold emblem flying to the right. The pole is perfectly vertical.", BANNER)

# ---- 3. gate
add('G1_gatehouse.png', 'Main gatehouse', 'sprite', (1, 1), "a castle main gatehouse",
    "A wide stone gate building with two half-round towers on its left and right, a large pointed arch in the exact centre with a dark empty passage inside (no gate, no portcullis, no doors), battlements along the top, a carved stone shield above the arch, a crimson banner on each tower, an iron wall torch with a small flame on each side of the arch, a few arrow slits. Perfectly symmetric. Only the front face is visible.", STONE + "; " + BANNER + "; " + IRON + "; " + FIRE + "; passage #0c0e16")
add('G2_portcullis.png', 'Portcullis', 'sprite', (5, 6), "an iron portcullis",
    "A closed iron portcullis grid: thick dark vertical bars crossed by horizontal bars, pointed spikes at the bottom of each vertical bar, rectangular overall. All bars perfectly straight and parallel.", IRON)
add('G3A_gate_doors_closed.png', 'Gate doors closed', 'sprite', (5, 6), "a pair of closed castle gate doors",
    "Two heavy oak doors closed together, black iron straps and studs, two iron ring handles, together shaped like a pointed arch. Symmetric.", WOOD + "; " + IRON)
add('G3B_gate_doors_open.png', 'Gate doors open', 'sprite', (5, 6), "a pair of castle gate doors swung open",
    "The same oak and iron gate doors swung open inward, seen as two narrow door edges at the left and right, a dark stone passage between them, pointed arch outline. Symmetric.", WOOD + "; " + IRON + "; passage #0c0e16")

# ---- 4. stairs
add('S1_stairs_upper.png', 'Upper grand stairs (plain)', 'sprite', (4, 3), "a wide straight stone staircase going up away from the viewer",
    "Exactly 9 wide stone steps rising straight away from the viewer, seen from the front and slightly above: every step edge is a perfectly horizontal line, treads lit, risers in shade. A carved stone balustrade with square posts on both sides. Symmetric.", STONE + "; " + FLAG)
add('S2_stairs_lower_carpet.png', 'Lower grand stairs (red carpet)', 'sprite', (1, 1), "a wide straight stone staircase with a red carpet",
    "Exactly 11 wide stone steps rising straight away from the viewer, seen from the front and slightly above, every step edge perfectly horizontal. A wine-red carpet with gold edging runs straight up the exact centre over every step. Carved stone balustrades on both sides ending in square posts with small glowing lamps. Symmetric.", STONE + "; " + CARPET + "; lamp glow #ffd27a")
add('S3_balustrade.png', 'Stone balustrade', 'band', None, "a low carved stone balustrade railing",
    "Round stone balusters evenly spaced under a flat straight handrail, standing on a straight stone base. Straight-on front view.", STONE)
add('S4_gargoyle.png', 'Gargoyle statue', 'sprite', (5, 6), "a winged gargoyle statue",
    "A crouching winged gargoyle of dark iron-grey stone, wings half raised, facing the viewer, on a small square stone base. Symmetric.", IRON + "; " + STONE)
add('S5_knight_statue_sword.png', 'Knight statue (sword)', 'sprite', (1, 2), "a knight statue on a plinth",
    "A dark iron knight statue in full plate armour standing to attention, facing the viewer, both hands resting on the pommel of a sword planted point-down in front, on a square stone plinth. Symmetric.", IRON + "; " + STONE)

# ---- 5. palace
add('P1_palace_keep.png', 'Palace central keep', 'sprite', (4, 5), "the central keep of a royal fantasy palace",
    "A tall Gothic stone palace block: pinnacles and four slender spires with blue cone roofs, a huge crimson banner with a gold cross hanging down the centre, a grand arched double door of dark wood at the bottom centre with a small stone balcony above it, two knight statues in niches beside the door, rows of tall windows glowing warm yellow, blue slate roofs. Perfectly symmetric.", STONE + "; " + ROOF + "; " + BANNER + "; " + GLOW + "; " + WOOD)
add('P2_palace_wing_a.png', 'Palace wing A', 'sprite', (3, 4), "a straight section of a royal palace wing",
    "A rectangular stone palace wall section, two storeys of tall Gothic windows glowing warm yellow, stone buttresses between the windows, a steep blue slate roof with two dormer windows and a gold ridge on top, small crimson banners between windows. The wall and roof run straight to the left and right edges of the building so several copies can be placed side by side.", STONE + "; " + ROOF + "; " + GLOW + "; " + BANNER)
add('P3_palace_wing_b.png', 'Palace wing B', 'sprite', (3, 4), "a straight section of a royal palace wing (variant)",
    "The same palace wing style, same height, same roof line, but with a central projecting bay window and a stone chimney, so it looks different from the plain wing. The wall and roof run straight to the left and right edges of the building.", STONE + "; " + ROOF + "; " + GLOW)
add('P4_palace_pavilion.png', 'Palace corner pavilion', 'sprite', (1, 2), "a square palace corner pavilion",
    "A tall square stone pavilion with three storeys of glowing Gothic windows, a steep pyramid blue slate roof with a gold finial, one crimson banner. Symmetric.", STONE + "; " + ROOF + "; " + GLOW + "; " + BANNER)
add('P5_palace_skyline_left_half.png', 'Palace skyline (left half)', 'sprite', (21, 9), "the LEFT HALF of a palace roof skyline",
    "A row of 5 slender stone towers of different heights with steep blue cone roofs and gold finials, joined by blue slate roofs, a few warm glowing windows. The towers get taller toward the RIGHT; the tallest centre spire is cut exactly in half by the right edge of the image (we mirror it to make the full skyline). The bottom of the buildings is one straight horizontal line.", STONE + "; " + ROOF + "; " + GLOW)
add('P6_palace_giant_banner.png', 'Giant palace banner', 'sprite', (9, 16), "a very long royal banner",
    "A very long crimson banner hanging straight down from a gold rod, a large gold cross-and-fleur emblem in the upper middle, gold border, pointed bottom. Hanging still, perfectly symmetric.", BANNER)

# ---- 6. buildings
add('B1_cathedral.png', 'Cathedral', 'sprite', (3, 4), "the front of a Gothic cathedral",
    "A large circular stained-glass rose window above a pointed-arch double door, two tall spires with blue slate roofs, stone pinnacles, tall lit pointed windows, a steep blue roof behind, two crimson banners, a short stone entrance stair. Perfectly symmetric.", STONE + "; " + ROOF + "; " + GLOW + "; " + BANNER + "; " + WOOD)
add('B2_chapel.png', 'Chapel', 'sprite', (2, 3), "the front of a small Gothic chapel",
    "A pointed-arch wooden door, a small round stained-glass window above it, a steep blue slate roof, a slender bell turret with a cross on top. Symmetric.", STONE + "; " + ROOF + "; " + GLOW + "; " + WOOD)
add('B3_domed_hall.png', 'Domed hall', 'sprite', (1, 1), "the front of a grand domed hall",
    "A central dark-blue tiled dome with a small lantern cupola and gold finial, a ring of glowing arched windows under the dome, a stone facade with columns, two small side wings with blue roofs, a wide entrance stair in the centre. Perfectly symmetric.", STONE + "; " + ROOF + "; " + GLOW + "; " + WOOD)
add('B4_manor.png', 'Manor house', 'sprite', (4, 3), "the front of a stone manor house",
    "Two storeys of stone walls with rows of warm glowing windows, a blue slate hipped roof with dormers and two chimneys, a central arched oak door with a small porch, two small crimson banners. Symmetric.", STONE + "; " + ROOF + "; " + GLOW + "; " + WOOD + "; " + BANNER)
add('B5_colonnade.png', 'Arched colonnade', 'sprite', (2, 1), "a stone colonnade wall",
    "A straight stone wall with exactly 3 equal tall round arches with dark shadowed insides, slender columns between the arches, a row of battlements on top. Symmetric.", STONE + "; arch insides #0c0e16, #1a1d2a")
add('B6_fountain_dry.png', 'Central fountain (no water)', 'sprite', (1, 1), "an ornate tiered stone fountain WITHOUT water",
    "An octagonal stone basin with a carved rim, a central column holding two stacked round bowls and a small finial on top, small flower planters at the basin corners. The basin and bowls are EMPTY and dry: absolutely no water, no spray. Seen from the front and slightly above. Symmetric.", STONE + "; flowers #b3262e, #e0b040, #3f8f4f")

# ---- 7. floors and surfaces
add('T8_courtyard_cobbles.png', 'Courtyard cobbles', 'tile', (1, 1), "castle courtyard cobblestones seen from directly above",
    "Irregular rounded cobblestones of mixed sizes packed together, dark joints, a little moss in some joints, a few cracked stones. Calm and even.", COBBLE)
add('T9_terrace_paving.png', 'Palace terrace paving', 'tile', (1, 1), "large square palace paving slabs seen from directly above",
    "Large neatly cut square and rectangular flagstones in a regular pattern, thin precise joints, very light wear. Clean and pale.", FLAG)
add('T10_road_cobbles.png', 'Causeway road', 'tile', (1, 1), "a cobbled road surface seen from directly above",
    "Rounded cobbles, slightly warmer tone, a smoother wheel-worn band running straight from top to bottom through the middle, small gravel between stones.", COBBLE)
add('T11_lawn.png', 'Formal lawn', 'tile', (1, 1), "a short trimmed garden lawn seen from directly above",
    "Fine short grass blades, very even, low contrast, a couple of tiny clover patches. No flowers, no stones.", GRASS)
add('T12_dirt_yard.png', 'Dirt yard', 'tile', (1, 1), "a packed earth yard seen from directly above",
    "Packed tan earth with scattered small pebbles, faint footprints and cart ruts, two or three tiny dry grass tufts.", DIRT)
add('T13_bridge_arch.png', 'Stone bridge arch', 'sprite', (3, 2), "the front face of one stone bridge arch",
    "A stone wall section with ONE large round arch in the centre: a ring of wedge-shaped stones with a keystone at the top, a short pier on each side with dark wet algae at the bottom. The OPENING under the arch is filled with flat magenta #FF00FF (see-through). The stone wall runs to the left and right edges of the image so arches can be placed side by side.", STONE + "; " + WET)
add('T14_cliff_face.png', 'Cliff face', 'band', None, "a vertical rock cliff face under a grassy edge",
    "Grey-blue vertical rock columns and ledges, a strip of green grass with a few hanging tufts along the top edge of the strip, darker wet rock along the bottom edge. The strip is tall: about one third of the image height.", ROCK + "; " + GRASS + "; " + WET)
add('T15_pool_rim.png', 'Pool rim', 'band', None, "the carved stone edge of an ornamental pool",
    "A rounded stone lip with a lit top surface and a short shaded front face, perfectly straight.", FLAG)

# ---- 8. props
add('F1_knight_statue_shield.png', 'Knight statue (shield)', 'sprite', (1, 2), "a knight statue on a pedestal",
    "A dark iron knight statue in plate armour holding a tall shield in front and a sword, facing the viewer, on a square stone pedestal with a carved crest.", IRON + "; " + STONE)
add('F2_lantern_pillar.png', 'Lantern pillar', 'sprite', (2, 5), "a stone lantern pillar",
    "A square carved stone pillar with a black iron lantern on top, the lantern's glass panes glowing warm yellow. Symmetric.", STONE + "; " + IRON + "; " + GLOW)
add('F3_brazier.png', 'Brazier (no flame)', 'sprite', (2, 5), "an iron fire brazier on a tripod",
    "A round black iron bowl filled with glowing orange embers on a tall three-legged iron tripod. NO flames above the bowl.", IRON + "; embers #3a1a10, #ff6a1a, #ffb52e")
add('F4_wall_torch.png', 'Wall torch (no flame)', 'sprite', (1, 2), "an iron wall torch bracket",
    "A black iron wall bracket holding a wooden torch with a charred top, as mounted on a wall. NO flame. Perfectly vertical.", IRON + "; " + WOOD + "; charred #1a1410")
add('F5_wall_banner.png', 'Wall banner', 'sprite', (1, 2), "a castle wall banner",
    "A crimson banner hanging straight down from a short iron rod, gold cross emblem, thin gold border, swallow-tail bottom. Hanging still, symmetric.", BANNER)
add('F6_sluice_grate.png', 'Sluice grate', 'sprite', (5, 6), "a small iron water grate in a stone arch",
    "A round-topped stone arch with straight vertical iron bars inside, dark stains below. Symmetric.", IRON + "; " + STONE + "; stain #2a3a4a")
add('F7A_arrow_slit.png', 'Arrow slit decal', 'sprite', (1, 2), "a single arrow slit in a stone wall",
    "A narrow vertical arrow slit opening framed by a few stone blocks, dark inside. Only this small wall detail.", STONE + "; opening #0c0e16")
add('F7B_drain_spout.png', 'Drain spout decal', 'sprite', (1, 1), "a stone drain spout sticking out of a wall",
    "A small carved stone spout with a dark stain running down below it. Only this detail.", STONE + "; stain #2a3a4a")
add('F7C_ivy_patch.png', 'Ivy patch decal', 'sprite', (1, 1), "a patch of ivy hanging on a wall",
    "A loose cluster of dark green ivy leaves and thin vines hanging downward. Only the ivy, no wall.", "ivy #1f3326, #24402a, #335a33, #4f7f3f")
add('F7D_barred_window.png', 'Barred window decal', 'sprite', (1, 1), "a small barred castle window",
    "A small square stone window frame with three vertical iron bars, dark inside. Symmetric.", STONE + "; " + IRON + "; inside #0c0e16")

# ---- 9. effects (single images; animated in code)
add('E1_waterfall.png', 'Waterfall', 'sprite', (2, 3), "a waterfall",
    "White-blue water falling straight down in three wide streams that merge, foamy and bright, falling perfectly vertically, slightly wider at the bottom. Only the falling water: no rocks, no cliff, no pool.", WATER)
add('E2_splash.png', 'Waterfall splash', 'sprite', (3, 1), "white foam and spray where a waterfall hits a lake",
    "A wide low cloud of white churning foam, spray droplets and bubbles spreading out to the left and right. Only the foam: no water surface, no rocks.", WATER)
add('E4_fountain_water.png', 'Fountain water', 'sprite', (1, 1), "the water of a tiered fountain",
    "Only the water of a two-bowl tiered fountain: a small jet at the top, thin curtains of water spilling over the edges of two round bowls, and a calm octagonal pool of blue water at the bottom. No stone at all.", WATER)
add('E5_flame.png', 'Torch flame', 'sprite', (1, 2), "a small torch flame",
    "A single bright flame: white-yellow core, orange body, red tips, two or three sparks, pointing straight up. Only the flame.", FIRE)
add('E7A_seagull_wings_up.png', 'Seagull (wings up)', 'sprite', (3, 2), "a flying seagull",
    "A small white seagull with grey wingtips flying to the right, side view, wings raised up.", "white #f4f6f8, #d8dde3; grey #4a5060; beak #e0a030")
add('E7B_seagull_wings_down.png', 'Seagull (wings down)', 'sprite', (3, 2), "a flying seagull",
    "A small white seagull with grey wingtips flying to the right, side view, wings pushed down.", "white #f4f6f8, #d8dde3; grey #4a5060; beak #e0a030")

# ---- 10. terrain
add('R1A_rock_outcrop_large.png', 'Rock outcrop (large)', 'sprite', (2, 1), "a large mossy rock outcrop",
    "A wide mound of grey-blue rocks with patches of moss and grass on top, flat bottom edge (it will sit in water). Only the rock.", ROCK + "; " + GRASS)
add('R1B_rock_outcrop_small.png', 'Rock outcrop (small)', 'sprite', (3, 2), "a small mossy rock outcrop",
    "A small cluster of grey-blue rocks with a little moss on top, flat bottom edge. Only the rock.", ROCK + "; " + GRASS)
add('R2_shore_edge.png', 'Shore edge', 'band', None, "a lake shoreline seen from above",
    "Green grass along the top edge of the strip turning into pebbles and wet sand along the bottom edge.", GRASS + "; pebbles #7a7470, #a39a90; wet sand #8f7a5e, #b8a080")
add('T5A_corner_outer_90.png', 'Battlement corner (outer 90°)', 'sprite', (1, 1), "one square stone battlement block (merlon) for the outside corner of a castle wall",
    "ONE square stone merlon block seen exactly from the FRONT, like a flat building drawing: we see its front face as an upright square of stone blocks and only a thin lit strip of its top face above it. "
    "It must NOT look like a 3D cube turned at 45 degrees: no diamond shape, no two side faces, no isometric corner pointing at the viewer. Left and right edges are perfectly vertical. Only this one block.", STONE)
add('T5B_corner_inner_90.png', 'Battlement corner (inner 90°)', 'sprite', (2, 1), "a short piece of castle battlement that turns an inside corner",
    "Seen exactly from the FRONT, like a flat building drawing: two stone merlons side by side on a short straight stone base; the RIGHT merlon is set slightly further back (higher up in the image) to show the wall turning away from us. "
    "It must NOT look isometric: no diamond or L shape on the floor, no faces turned at 45 degrees, all vertical edges perfectly vertical, the base perfectly horizontal.", STONE)

# ---------------------------------------------------------------- prompt text
NO_GROUND_RULE = {'T2', 'T14', 'R2', 'E1', 'E2', 'E4', 'R1A', 'R1B', 'F7C'}
def build(r):
    save, title, kind, shp, subject, details, colours = r
    key = save.split('_')[0]
    ground = '' if (kind == 'band' or key in NO_GROUND_RULE) else 'a floor, ground, grass or water under the object; '
    dont = DONT_SPRITE.replace('{GROUND}', ground)
    view = VIEW_SIDE if key.startswith('E7') else VIEW
    if kind == 'tile':
        shape = shape_of(*shp)
        parts = [f"Create a seamless pixel-art texture: {subject}.", f"DETAILS: {details}", ref_block(key, kind), f"COLOURS (use only these): {colours}.",
                 STYLE, TILE, f"IMAGE SHAPE: {shape}.", DONT_TILE]
    elif kind == 'band':
        shape = '16:9 wide'
        parts = [f"Create a pixel-art game asset: {subject}.", f"DETAILS: {details}", ref_block(key, kind), f"COLOURS (use only these): {colours}.",
                 STYLE, BAND, f"IMAGE SHAPE: {shape}.", dont.replace("anything cut off by the image edge", "anything cut off at the top or bottom of the strip")]
    else:
        shape = shape_of(*shp)
        parts = [f"Create ONE pixel-art game asset: {subject}.", f"DETAILS: {details}", ref_block(key, kind), f"COLOURS (use only these): {colours}.",
                 STYLE, view, BG, f"IMAGE SHAPE: {shape}.", dont]
    return shape, "\n\n".join(parts)

# ---------------------------------------------------------------- one-sheet workbook
F = 'Arial'
wb = Workbook(); ws = wb.active; ws.title = 'Prompts'
thin = Side(style='thin', color='9AA0B0'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
ws.column_dimensions['A'].width = 6; ws.column_dimensions['B'].width = 34; ws.column_dimensions['C'].width = 13
ws.column_dimensions['D'].width = 120; ws.column_dimensions['E'].width = 16; ws.column_dimensions['F'].width = 10
import json as _json
_SP = os.path.join(OUT, 'status.json')
STATUS = _json.load(open(_SP)) if os.path.exists(_SP) else {}   # SAVE AS -> "Done" / "Redo" (kept between rebuilds)
PEOPLE = ['Gokhul', 'Bhuvanesh', 'Surendar']; PCOL = {'Gokhul': 'DDEBF7', 'Bhuvanesh': 'FCE4D6', 'Surendar': 'E2EFDA'}
def who(n): return PEOPLE[min((n - 1) // 21, 2)]   # 1-21 Gokhul, 22-42 Bhuvanesh, 43-65 Surendar
ws.merge_cells('A1:F1'); ws['A1'] = 'CASTLE ASSETS: copy the prompt, generate, save with the file name. That is all.'
ws['A1'].font = Font(name=F, bold=True, size=15, color='FFFFFF'); ws['A1'].fill = PatternFill('solid', fgColor='1C2036')
ws['A1'].alignment = Alignment(vertical='center', indent=1); ws.row_dimensions[1].height = 30
steps = ('0) With EVERY prompt, attach these 3 images in this order: 1_castle_grid.png, 2_legend.png, 3_castle_labels.png (in the same folder as this file).   '
         '1) Copy the whole cell in column "PROMPT" and paste it into the image generator as it is (do not shorten it).   '
         '2) Set the image shape shown in column "SHAPE" if the tool has that option.   '
         '3) Save the image as PNG with the exact name in column "SAVE AS" and put it in the shared folder.   '
         '4) Only do the rows with YOUR name in column "ASSIGNED TO". Mark "Done" in the last column. Rows marked "Redo" (red) need a new image with the updated prompt.   If an image looks tilted, has a person in it, or has more than one object: generate it again.')
ws.merge_cells('A2:F2'); ws['A2'] = steps; ws['A2'].font = Font(name=F, size=11); ws['A2'].alignment = Alignment(wrap_text=True, vertical='center', indent=1)
ws['A2'].fill = PatternFill('solid', fgColor='FFF2B3'); ws.row_dimensions[2].height = 80
hdr = ['No.', 'SAVE AS (exact file name)', 'SHAPE / LOOK AT CELL', 'PROMPT (copy everything in the cell) + attach the 3 images', 'ASSIGNED TO', 'Done']
for i, h in enumerate(hdr, 1):
    c = ws.cell(3, i, h); c.font = Font(name=F, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2B3249')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
ws.row_dimensions[3].height = 24
txt = ['CASTLE ASSET PROMPTS', '=' * 60,
       'With EVERY prompt attach 3 images in this order: 1_castle_grid.png, 2_legend.png, 3_castle_labels.png',
       'For each item: copy the PROMPT, generate, save as PNG with the SAVE AS name.',
       'If an image is tilted, has a person, or has more than one object: generate it again.', '']
for n, r in enumerate(L, 1):
    shape, p = build(r); row = n + 3
    _cs = cells_for(r[0].split('_')[0]); vals = [n, r[0], shape + ('\n\nCell ' + _cs[0] if _cs else '\n\n(style only)'), p, who(n), STATUS.get(r[0])]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v); c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='top', horizontal='center' if i in (1, 3, 5, 6) else 'left')
        c.font = Font(name=F, size=10, bold=(i == 2))
    ws.cell(row, 2).value = f"{r[0]}\n\n{r[1]}"
    ws.row_dimensions[row].height = 409
    txt += [f'#{n}  ASSIGNED TO: {who(n).upper()}' + (f'   >>> {STATUS[r[0]].upper()} <<<' if r[0] in STATUS else '') + '   SAVE AS: {r[0]}   ({r[1]})   SHAPE: {shape}', '-' * 60, p, '', '']
last = len(L) + 3
dv = DataValidation(type='list', formula1='"Done,Redo"', allow_blank=True); ws.add_data_validation(dv); dv.add(f'F4:F{last}')
ws.conditional_formatting.add(f'A4:F{last}', FormulaRule(formula=[f'$F4="Done"'], fill=PatternFill('solid', fgColor='C6EFCE')))
ws.conditional_formatting.add(f'A4:F{last}', FormulaRule(formula=[f'$F4="Redo"'], fill=PatternFill('solid', fgColor='FFC7CE')))
for _r in range(4, last + 1):
    _c = ws.cell(_r, 5); _c.font = Font(name=F, size=11, bold=True); _c.fill = PatternFill('solid', fgColor=PCOL[_c.value]); _c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
for _i, _p in enumerate(PEOPLE):
    ws.cell(last + 3 + _i, 2, f'{_p} done:').font = Font(name=F, bold=True)
    ws.cell(last + 3 + _i, 3, f'=COUNTIFS(E4:E{last},"{_p}",F4:F{last},"Done")&" / "&COUNTIF(E4:E{last},"{_p}")').font = Font(name=F, bold=True)
ws.cell(last + 2, 2, 'Images done:').font = Font(name=F, bold=True)
ws.cell(last + 2, 3, f'=COUNTIF(F4:F{last},"Done")&" / "&COUNTA(B4:B{last})').font = Font(name=F, bold=True)
ws.freeze_panes = 'A4'
import shutil
for src, dst in [('asset_map_grid_only.png', '1_castle_grid.png'), ('asset_map_legend.png', '2_legend.png'), ('asset_map_full.png', '3_castle_labels.png')]:
    shutil.copy(os.path.join(ROOT, 'docs/asset_map', src), os.path.join(OUT, dst))
wb.save(os.path.join(OUT, 'CASTLE_PROMPTS.xlsx'))
open(os.path.join(OUT, 'CASTLE_PROMPTS.txt'), 'w').write('\n'.join(txt))
print(len(L), 'prompts')
